"""Loads the QRC rulebook (Quick_Reference_Card_-_SEA_T_T.pptx) into rule_chunks.

Chunking: one chunk per rule block, semantic, never fixed-size (MASTER §5,
DEV_A §3.3). Text is extracted at runtime from the pptx (not hand-transcribed)
so the slide-14 footer stays byte-identical to the source — the synthesis
agent must never paraphrase it.
"""

from __future__ import annotations

import logging

from pptx import Presentation

log = logging.getLogger(__name__)


def _slide_paragraphs(slide) -> list[str]:
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = "".join(run.text for run in para.runs).strip()
                if t:
                    texts.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    texts.append(" | ".join(cells))
    return texts


def _join(paras: list[str]) -> str:
    return "\n".join(paras)


def _split_before(paras: list[str], marker_prefix: str) -> tuple[list[str], list[str]]:
    """Splits paras into (before, from-marker-onward), matching by prefix."""
    for i, p in enumerate(paras):
        if p.startswith(marker_prefix):
            return paras[:i], paras[i:]
    return paras, []


def _exclude(paras: list[str], exact_values: set[str]) -> list[str]:
    return [p for p in paras if p not in exact_values]


def build_chunks(pptx_path: str) -> list[dict]:
    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    chunks: list[dict] = []

    def add(slide_no: int, title: str, source_db: str, text: str, chunk_type: str = "narrative"):
        idx = sum(1 for c in chunks if c["slide_number"] == slide_no) + 1
        chunks.append({
            "chunk_id": f"qrc-{slide_no}-{idx}",
            "slide_number": slide_no,
            "title": title,
            "source_db": source_db,
            "chunk_type": chunk_type,
            "text": f"[Slide {slide_no} — {title}]\n{text}",
        })

    # ---- Slide 2 — Assigning DCCS Request ----
    p2 = _exclude(_slide_paragraphs(slides[1]), {"Assigning DCCS Request"})
    add(2, "Assigning DCCS Request", "GENERAL", _join(p2))

    # ---- Slide 3 — engagement-details QC / referral work ----
    p3 = _slide_paragraphs(slides[2])
    p3 = _exclude(p3, {"Perform Quality Check – Engagement Details", "Referral Work"})
    qc_part, referral_part = _split_before(p3, "Check if request is for referral work")
    add(3, "Perform Quality Check – Engagement Details", "GENERAL", _join(qc_part))
    add(3, "Referral Work", "GENERAL", _join(referral_part))

    # ---- Slide 4 — relevant parties QC / DESC inconsistency ----
    p4 = _slide_paragraphs(slides[3])
    p4 = _exclude(p4, {"Perform Quality Check – Relevant Parties"})
    qc_parties, desc_inconsistency = _split_before(p4, "Make sure key information")
    add(4, "Perform Quality Check – Relevant Parties", "DESC", _join(qc_parties))
    add(4, "DESC Is Authoritative / Inconsistency Handling", "DESC", _join(desc_inconsistency))

    # ---- Slide 5 — DESC search + designation types ----
    p5 = _slide_paragraphs(slides[4])
    p5 = _exclude(p5, {
        "Perform Search to Conflicts Databases – DESC",
        "To Check Entity’s DESC Tree",
        "To Check DESC LCSP/RP",
    })
    add(5, "Perform Search to Conflicts Databases – DESC", "DESC", _join(p5))

    # ---- Slide 6 — WBS business-unit / status classification ----
    p6 = _slide_paragraphs(slides[5])
    p6 = _exclude(p6, {"Perform Search to Conflicts Databases – WBS", "Sample Case"})
    bu_part, status_part = _split_before(p6, "For Status, please indicate")
    add(6, "WBS Business-Unit Classification", "WBS", _join(bu_part))
    add(6, "WBS Status Classification", "WBS", _join(status_part))

    # ---- Slide 7 — One Window intro + Strategic Client ----
    p7 = _slide_paragraphs(slides[6])
    p7 = _exclude(p7, {"Perform Search to Conflicts Databases – New One Window (Sharepoint)"})
    add(7, "One Window — Strategic Client", "ONE_WINDOW", _join(p7))

    # ---- Slide 8 — Directorship / Sanction / Open Opportunities ----
    p8 = _slide_paragraphs(slides[7])
    p8 = _exclude(p8, {"Perform Search to Conflicts Databases – New One Window (Sharepoint)"})
    directorship, rest = _split_before(p8, "Sanction Match")
    sanction, open_opp = _split_before(rest, "Open Opportunities Listing")
    add(8, "One Window — Directorship Listing", "ONE_WINDOW", _join(directorship))
    add(8, "One Window — Sanction Match", "ONE_WINDOW", _join(sanction))
    add(8, "One Window — Open Opportunities Listing", "ONE_WINDOW", _join(open_opp))

    # ---- Slide 9 — COT search + cutoff ----
    p9 = _slide_paragraphs(slides[8])
    p9 = _exclude(p9, {"Perform Search to Conflicts Databases – COT"})
    add(9, "Perform Search to Conflicts Databases – COT", "COT", _join(p9))

    # ---- Slide 10 — DCCS search result relevance tests ----
    p10 = _slide_paragraphs(slides[9])
    p10 = _exclude(p10, {"Perform Search to Conflicts Databases – DCCS Search Result"})
    add(10, "DCCS Search Result Relevance Tests", "DCCS_SEARCH", _join(p10))

    # ---- Slide 11 — cross-border matrix, kept whole ----
    p11 = _slide_paragraphs(slides[10])
    p11 = _exclude(p11, {"Perform Cross-Border Conflict Check Request to other Member Firms"})
    add(11, "Cross-Border Matrix", "CROSS_BORDER", _join(p11), chunk_type="decision_table")

    # ---- Slide 12 — sending a cross-border request ----
    p12 = _slide_paragraphs(slides[11])
    p12 = _exclude(p12, {"Perform Cross-Border Conflict Check Request to other Member Firms"})
    add(12, "Sending a Cross-Border Request", "CROSS_BORDER", _join(p12))

    # ---- Slide 13 — collate relationships, final result options ----
    p13 = _slide_paragraphs(slides[12])
    p13 = _exclude(p13, {"Collate all Relationships identified"})
    add(13, "Collate Relationships / Final Result", "GENERAL", _join(p13))

    # ---- Slide 14 — footer, verbatim ----
    p14 = _slide_paragraphs(slides[13])
    p14 = _exclude(p14, {
        "Additional Footer for All SEA T&T Request",
        "Application of Footer in DCCS Result",
    })
    reminders_part, personal_part = _split_before(p14, "Personal Conflicts:")
    add(14, "Reminders", "FOOTER", _join(reminders_part), chunk_type="decision_table")
    add(14, "Personal Conflicts", "FOOTER", _join(personal_part), chunk_type="decision_table")

    if len(chunks) < 10:
        log.warning("Only %d rule chunks parsed from QRC — check slide indices", len(chunks))

    return chunks


def ingest_qrc(session, pptx_path: str) -> int:
    from sqlalchemy import text as sql

    chunks = build_chunks(pptx_path)
    session.execute(sql("TRUNCATE TABLE rule_chunks RESTART IDENTITY CASCADE"))
    for chunk in chunks:
        session.execute(
            sql(
                """
                INSERT INTO rule_chunks (chunk_id, slide_number, title, source_db, chunk_type, text, search_vec)
                VALUES (:chunk_id, :slide_number, :title, :source_db, :chunk_type, :text,
                        to_tsvector('english', :title || ' ' || :text))
                """
            ),
            chunk,
        )
    session.commit()
    log.info("Parsed %d QRC rule chunks", len(chunks))
    return len(chunks)
